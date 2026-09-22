#!/usr/bin/env python3
"""python-pptx native builder.

Input:  step1_facts_sheet.json, step2_story_outline.json, step3_visual_spec.json
from workspace/outs/ (step3 carries top-level theme/color_palette per its schema;
a legacy nested visual_spec block is still honored if present).

Output: workspace/outs/final_deck.pptx

Usage:
    python3 tools/pptx_compiler.py    # run from the repo root

Exit codes:
    0   deck written
    1   missing inputs or invalid palette (failures are loud, never silent)
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
except ImportError:
    sys.exit("pptx not installed. Run: pip install python-pptx")

OUTPUTS_DIR = Path("workspace") / "outs"


def parse_hex(hex_str):
    """Convert '#RRGGBB' string to RGBColor (0-255 ints)."""
    hex_str = hex_str.lstrip("#")
    return RGBColor(
        int(hex_str[0:2], 16),
        int(hex_str[2:4], 16),
        int(hex_str[4:6], 16),
    )


def load_json_or_exit(path: Path):
    if not path.exists():
        print(f"[error] {path} not found")
        sys.exit(1)
    with open(path) as f:
        return json.load(f)


def load_spec_from_outputs():
    """Read the three pipeline artifacts from workspace/outs/ and merge into one spec."""
    if not OUTPUTS_DIR.exists():
        print("[error] workspace/outs/ directory not found")
        sys.exit(1)

    s1 = load_json_or_exit(OUTPUTS_DIR / "step1_facts_sheet.json")
    s2 = load_json_or_exit(OUTPUTS_DIR / "step2_story_outline.json")
    s3 = load_json_or_exit(OUTPUTS_DIR / "step3_visual_spec.json")

    # step3's schema puts theme/color_palette at the TOP LEVEL; a nested
    # visual_spec block is legacy. Normalize both into visual_spec.
    visual_spec = s3.get("visual_spec") or {
        k: v for k, v in s3.items() if k in ("theme", "color_palette")
    }

    return {
        "facts_sheet": s3.get("facts_sheet") or s1,
        "story_outline": s3.get("story_outline") or s2,
        "visual_spec": visual_spec,
        "meta": s3.get("meta", {}),
    }


def build(spec: dict, output: str) -> None:
    theme = spec.get("visual_spec", {}).get("theme", "Modern_Tech_Dark")
    palette = spec.get("visual_spec", {}).get("color_palette", {})

    prs = Presentation()
    # Use Title + Content layout (index 1)
    slide_layout = prs.slide_layouts[1]

    slides_data = spec.get("story_outline", {}).get("slides", [])
    if not slides_data:
        # Fallback: single title slide using first fact if available
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = spec.get("meta", {}).get("target_audience", "Deck")
        facts = spec.get("facts_sheet", {}).get("facts", [])
        if facts:
            content = slide.placeholders[1]
            content.text = facts[0].get("claim", "")
        print(f"Warning: no slide data; created title-only slide in {output}")
    else:
        for i, slide_info in enumerate(slides_data):
            slide = prs.slides.add_slide(slide_layout)

            # Title
            title = slide.shapes.title
            title.text = slide_info.get("title", f"Slide {slide_info.get('slide_no', i+1)}")

            # Content placeholders
            content = slide.placeholders[1]
            core_msg = slide_info.get("core_message", "")
            evidence_ids = slide_info.get("evidence_ids", [])
            facts = []

            for eid in evidence_ids:
                for f in spec.get("facts_sheet", {}).get("facts", []):
                    if f.get("id") == eid:
                        facts.append(f"• {f['claim']}")

            tf = content.text_frame
            tf.text = core_msg if core_msg else ""
            for b in facts:
                p = tf.add_paragraph()
                p.text = b

            # Apply theme colors if palette available. A bad palette value
            # raises loudly instead of being silently swallowed.
            if palette:
                bg = slide.background
                fill = bg.fill
                fill.solid()
                fill.fore_color.rgb = parse_hex(palette.get("primary", "#0A192F"))

    prs.save(output)
    print(f"Saved {output} with {len(slides_data)} slides")


if __name__ == "__main__":
    argparse.ArgumentParser(description="Python-pptx deck builder").parse_args([])
    spec = load_spec_from_outputs()
    build(spec, str(OUTPUTS_DIR / "final_deck.pptx"))
